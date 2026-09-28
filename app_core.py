# -*- coding: utf-8 -*-
"""
app_core.py — ядро обработки микробиологических данных.

Модуль не зависит от графического интерфейса: здесь загрузка и предобработка
данных из Excel, расчёт статистики по антибиотикограммам, формирование
аналитических листов и сохранение результата (в т.ч. с современным
оформлением через openpyxl).
"""

import os
from datetime import datetime

import pandas as pd
from openpyxl.utils import get_column_letter

# ----------------------------------------------------------------------
# 1. Загрузка и предобработка данных (расширенная)
# ----------------------------------------------------------------------

# Ключевые слова, по которым распознаём служебные столбцы (без жёстких названий)
_HEADER_KEYWORDS = ('заказ', 'история болезни', 'дата', 'материал', 'фио',
                    'patient', 'date', 'material', 'fio')

# Варианты названий ключевых столбцов (синонимы, регистронезависимо)
_ORGANISM_KEYS = ('микро', 'organism', 'возбудитель', 'идентификац', 'microbiolog',
                  'species', 'специес')
_DATE_KEYS = ('дат', 'date', 'число', 'дата посева', 'поступил')
_MATERIAL_KEYS = ('матер', 'биоматеріал', 'локализ', 'object', 'bid', 'вид материала',
                  'материал')
_PATIENT_KEYS = ('история', 'patient', 'номер', 'карта', 'закз', 'id', '№', 'п/п', 'фио', 'fio')


def _matches(col_norm, keys):
    return any(k in col_norm for k in keys)


def _norm(s):
    """Нормализация строки для нечувствительного сравнения."""
    return str(s).strip().lower().replace('\xa0', ' ')


def pick_sheet(xl, log_func=None):
    """Выбирает лист: сначала «TDSheet», затем первый лист с распознанными данными."""
    names = xl.sheet_names
    candidates = []
    for name in names:
        if _norm(name) == 'tdsheet':
            candidates.insert(0, name)
        else:
            candidates.append(name)

    for name in candidates:
        df_try = xl.parse(sheet_name=name, header=None)
        hdr = find_header_row(df_try)
        if hdr is not None:
            if log_func and name != candidates[0]:
                log_func(f"Лист '{name}': используется как источник данных "
                         f"(строка заголовков — {hdr + 1}).")
            return df_try, hdr
        if log_func:
            log_func(f"Лист '{name}' просмотрен — знакомых заголовков не найдено, пробуем следующий…")
    raise ValueError(
        "Не удалось распознать данные ни на одном листе. "
        "Убедитесь, что в файле есть таблица с колонками вроде «Дата», «Материал», «Микроорганизм»."
    )


def find_header_row(df_raw):
    """Ищет строку-заголовок без жёстких требований к оформлению.

    Подходит строка, первая ячейка которой содержит «Заказ»/«№»/«ID», либо
    строка с двумя и более известными заголовками («дата», «материал», «ФИО»…).
    Дополнительно строка признаётся заголовком, если в ней есть столбец про
    микроорганизм и ещё один распознанный столбец — это позволяет открывать
    файлы с нестандартными названиями колонок.
    """
    best_idx, best_score = None, 0
    limit = min(len(df_raw), 60)
    for i in range(limit):
        vals = [_norm(v) for v in df_raw.iloc[i].tolist() if pd.notna(v)]
        if len(vals) < 2:
            continue
        first = vals[0]
        score = 0
        if first.startswith('заказ') or first in ('№', 'no', 'no.', 'id'):
            score += 2
        matched = set()
        has_organism = False
        other_hits = 0
        for v in vals:
            if _matches(v, _ORGANISM_KEYS):
                has_organism = True
            for kw in _HEADER_KEYWORDS:
                if kw in v and kw not in matched:
                    matched.add(kw)
                    score += 1
                    break
            else:
                if (_matches(v, _DATE_KEYS) or _matches(v, _MATERIAL_KEYS)
                        or _matches(v, _PATIENT_KEYS)):
                    other_hits += 1
        # Комбинированный признак: микроорганизм + любой другой служебный столбец
        if has_organism and (other_hits >= 1 or len(matched) >= 1):
            score += 3
        if score >= 2 and score > best_score:
            best_idx, best_score = i, score
    return best_idx


def detect_ab_name_row(df_raw, organism_pos, header_row_idx, log_func=None):
    """Автоопределение строки с названиями антибиотиков.

    Не требуется, чтобы названия лежали в фиксированной 7-й строке: ищем строку
    (в пределах ~25 строк от начала), содержащую максимум текстовых значений
    в области столбцов антибиотиков.
    """
    ncols = df_raw.shape[1]
    start = organism_pos + 1
    if start >= ncols:
        return header_row_idx

    def count_names(r):
        cnt = 0
        for j in range(start, ncols):
            v = df_raw.iat[r, j]
            if pd.notna(v):
                s = str(v).strip()
                if s and not s.isdigit() and _norm(s) not in ('nan', 'none', 'rsi', 's', 'i', 'r', 'нд'):
                    cnt += 1
        return cnt

    best_r, best_c = None, 0
    for r in range(0, min(len(df_raw), max(header_row_idx + 2, 26))):
        c = count_names(r)
        if c > best_c:
            best_r, best_c = r, c
    if best_c >= 3 and best_r is not None:
        if log_func:
            log_func(f"Строка с названиями антибиотиков определена автоматически: №{best_r + 1}")
        return best_r
    if log_func:
        log_func("Не удалось уверенно найти строку с названиями антибиотиков — "
                 "используем строку сразу под заголовками.")
    return min(header_row_idx + 1, len(df_raw) - 1)


def load_data(filepath, log_func=None):
    if log_func:
        log_func("Чтение файла…")
    try:
        xl = pd.ExcelFile(filepath)
    except ValueError as e:
        # Старый формат .xls под Windows может определяться как ZIP (xlsx)
        msg = str(e)
        if 'zip' in msg.lower():
            try:
                xl = pd.ExcelFile(filepath, engine='openpyxl')
            except Exception:
                xl = pd.ExcelFile(filepath, engine='xlrd')
        else:
            raise

    # Лист и строка заголовков определяются автоматически —
    # требования к имени листа и расположению таблицы смягчены
    df_raw, header_row_idx = pick_sheet(xl, log_func=log_func)
    if header_row_idx is None:
        raise ValueError(
            "Не найдена строка с заголовками таблицы. Ожидаются колонки вида "
            "«Заказ / История болезни», «Дата», «Материал», «Микроорганизм»."
        )
    if log_func:
        log_func(f"Строка заголовков найдена: №{header_row_idx + 1}")

    # Основные заголовки берём из найденной строки
    headers = df_raw.iloc[header_row_idx].values
    data = df_raw.iloc[header_row_idx + 1:].reset_index(drop=True)
    data.columns = headers

    # Переименовываем основные колонки (в т.ч. по синонимам — без точного совпадения)
    rename_map = {}
    taken = set()
    for col in data.columns:
        col_str = str(col).strip()
        c_low = _norm(col_str)
        if col_str in ('История болезни', 'patient_id'):
            rename_map[col] = 'patient_id'
            taken.add('patient_id')
        elif col_str in ('Дата', 'date'):
            rename_map[col] = 'date'
            taken.add('date')
        elif col_str in ('Материал', 'material_raw'):
            rename_map[col] = 'material_raw'
            taken.add('material_raw')
        elif col_str in ('Микроорганизм', 'organism'):
            rename_map[col] = 'organism'
            taken.add('organism')
        elif col_str in ('ФИО', 'fio'):
            rename_map[col] = 'fio'
        else:
            # мягкое сопоставление по ключевым словам (первое подходящее имя выигрывает)
            if 'organism' not in taken and _matches(c_low, _ORGANISM_KEYS):
                rename_map[col] = 'organism'
                taken.add('organism')
            elif 'date' not in taken and _matches(c_low, _DATE_KEYS):
                rename_map[col] = 'date'
                taken.add('date')
            elif 'material_raw' not in taken and _matches(c_low, _MATERIAL_KEYS):
                rename_map[col] = 'material_raw'
                taken.add('material_raw')
            elif 'patient_id' not in taken and _matches(c_low, _PATIENT_KEYS):
                rename_map[col] = 'patient_id'
                taken.add('patient_id')
    data = data.rename(columns=rename_map)

    # Поиск столбца с микроорганизмом (расширенный, без жёстких названий)
    if 'organism' not in data.columns:
        for col in data.columns:
            c_low = _norm(col)
            if ('микро' in c_low or 'organism' in c_low or 'возбудитель' in c_low
                    or 'идентификац' in c_low or 'вид' == c_low):
                data = data.rename(columns={col: 'organism'})
                break
        else:
            raise ValueError(
                "Не найден столбец с микроорганизмом. Добавьте колонку с названием "
                "вида (например, «Микроорганизм» / «Возбудитель»)."
            )

    # Поиск столбца с датой (необязателен — при отсутствии заполняется пропуском)
    if 'date' not in data.columns:
        for col in data.columns:
            if 'дат' in _norm(col):
                data = data.rename(columns={col: 'date'})
                break
        else:
            if log_func:
                log_func("Столбец с датой не найден — статистика по датам будет недоступна.")
            data['date'] = pd.NaT

    # Поиск столбца с материалом (необязателен — при отсутствии используется «прочее»)
    if 'material_raw' not in data.columns:
        for col in data.columns:
            c_low = _norm(col)
            if 'матер' in c_low or 'биоматеріал' in c_low or 'локализ' in c_low or 'object' in c_low:
                data = data.rename(columns={col: 'material_raw'})
                break
        else:
            if log_func:
                log_func("Столбец с материалом не найден — все пробы будут отнесены к категории «другие».")
            data['material_raw'] = 'прочее'

    # Поиск столбца с пациентом (добавлена отказоустойчивость)
    if 'patient_id' not in data.columns:
        for col in data.columns:
            col_str = str(col).lower()
            if 'история' in col_str or 'patient' in col_str or 'номер' in col_str:
                data = data.rename(columns={col: 'patient_id'})
                break
        else:
            data['patient_id'] = data.index.astype(str)

    data['date'] = pd.to_datetime(data['date'], errors='coerce', dayfirst=True)

    # Находим позицию столбца 'organism'
    organism_pos = None
    for i, col in enumerate(data.columns):
        if col == 'organism':
            organism_pos = i
            break
    if organism_pos is None:
        raise ValueError("Не удалось определить позицию столбца с микроорганизмом")

    # Строка с названиями антибиотиков определяется автоматически —
    # больше не требуется фиксированное расположение (7-я строка)
    ab_name_row = detect_ab_name_row(df_raw, organism_pos, header_row_idx, log_func=log_func)
    if log_func:
        log_func(f"Используем строку {ab_name_row + 1} для названий антибиотиков")
        sample_cells = []
        for j in range(organism_pos+1, min(organism_pos+11, len(df_raw.columns))):
            val = df_raw.iloc[ab_name_row, j] if ab_name_row < len(df_raw) else ''
            sample_cells.append(str(val) if pd.notna(val) else 'NaN')
        log_func(f"Столбцы {organism_pos+2}-{organism_pos+11} на строке {ab_name_row}: {sample_cells}")

    antibiotics = []
    ab_indices = []
    for j in range(organism_pos + 1, len(data.columns), 2):
        if j+1 >= len(data.columns):
            break
        ab_name_raw = df_raw.iloc[ab_name_row, j] if ab_name_row < len(df_raw) else ''
        ab_name = str(ab_name_raw).strip()
        if not ab_name or ab_name.lower() in ('nan', 'none', ''):
            ab_name_raw = df_raw.iloc[ab_name_row, j+1] if ab_name_row < len(df_raw) else ''
            ab_name = str(ab_name_raw).strip()
        if ab_name and ab_name.lower() not in ('nan', 'none', ''):
            ab_name_clean = ab_name
            if (ab_name_clean and not ab_name_clean.isdigit()
                and ab_name_clean not in ('RSI', 'S', 'R', 'I', 'НД', 'Unnamed', 'ФИО', 'material_raw')):
                antibiotics.append(ab_name_clean)
                ab_indices.append(j+1)  # индекс RSI (правый столбец)

    if log_func:
        log_func(f"Найдено антибиотиков: {len(antibiotics)}. Примеры: {antibiotics[:10]}")
        if len(antibiotics) == 0:
            log_func("ВНИМАНИЕ: Не найдено ни одного антибиотика. Проверьте правильность строки с названиями.")

    # Функция диагностичности
    def is_diagnostic(row):
        if pd.isna(row['organism']) or str(row['organism']).strip().lower() in ('', 'роста нет', 'нет роста'):
            return False
        for idx in ab_indices:
            val = row.iloc[idx]
            if pd.notna(val) and isinstance(val, str):
                if val.strip().upper() in ('S', 'R', 'I'):
                    return True
        return False

    data['diagnostic'] = data.apply(is_diagnostic, axis=1)
    data['is_positive'] = data['organism'].apply(
        lambda x: False if pd.isna(x) or str(x).strip().lower() in ('', 'роста нет', 'нет роста') else True
    )

    data = data.dropna(subset=['material_raw'])
    # Пустые строки (без микроорганизма и без антибиотикограммы) отбрасываем,
    # но наличие даты больше не является обязательным
    data = data[~(data['organism'].isna() & ~data['diagnostic'])]
    data['material_raw'] = data['material_raw'].astype(str).str.strip()
    data['patient_id'] = data['patient_id'].astype(str).str.strip()

    df_main = data[['patient_id', 'date', 'material_raw', 'organism', 'is_positive', 'diagnostic']].copy()

    # Сбор данных по антибиотикам для всех образцов (группировка только по микроорганизму)
    ab_data = compute_ab_data(data, antibiotics, ab_indices, group_cols=['organism'], log_func=log_func)

    if log_func:
        log_func(f"Загружено {len(df_main)} строк, антибиотиков с данными: {len(ab_data)}")

    # Возвращаем также полный data, список антибиотиков и индексы для дальнейшего использования
    return df_main, ab_data, data, antibiotics, ab_indices


def compute_ab_data(df, antibiotics, ab_indices, group_cols=None, log_func=None):
    """
    Вычисляет статистику по антибиотикам для заданного DataFrame.
    df должен содержать столбец 'organism' и столбцы с RSI-значениями по индексам ab_indices.
    Параметр group_cols задаёт список столбцов для группировки (по умолчанию ['organism']).
    Возвращает словарь ab_data {антибиотик: DataFrame с подсчётами}.
    """
    if group_cols is None:
        group_cols = ['organism']

    ab_records = {ab: [] for ab in antibiotics}
    for idx, row in df[df['is_positive']].iterrows():
        org = row['organism']
        if pd.isna(org):
            continue
        # Формируем ключ группировки
        group_key = tuple(row[col] for col in group_cols)
        for i, ab in enumerate(antibiotics):
            rsi_val = row.iloc[ab_indices[i]]
            if pd.notna(rsi_val) and isinstance(rsi_val, str):
                rsi = rsi_val.strip().upper()
                if rsi in ('S', 'I', 'R', 'НД'):
                    # Сохраняем словарь с группирующими колонками и результатом
                    record = {col: row[col] for col in group_cols}
                    record['result'] = rsi
                    ab_records[ab].append(record)

    ab_data = {}
    for ab, records in ab_records.items():
        if records:
            df_ab = pd.DataFrame(records)
            # Группируем по group_cols и 'result'
            counts = df_ab.groupby(group_cols + ['result']).size().unstack(fill_value=0)
            # Убедимся, что все столбцы присутствуют
            for col in ['S', 'I', 'R', 'НД']:
                if col not in counts.columns:
                    counts[col] = 0
            counts = counts[['S', 'I', 'R', 'НД']]
            counts['Total'] = counts.sum(axis=1)
            for col in ['S', 'I', 'R', 'НД']:
                counts[f'{col}_%'] = counts[col] / counts['Total'] * 100
            ab_data[ab] = counts
        else:
            ab_data[ab] = pd.DataFrame()
    return ab_data


# ----------------------------------------------------------------------
# 2. Категоризация материалов для листов 2 и 3
# ----------------------------------------------------------------------

def categorize_material(material):
    material_lower = material.lower()
    if 'моча' in material_lower:
        return 'моча', 'моча'
    elif 'мокрота' in material_lower or 'аспират' in material_lower or 'бронхо' in material_lower:
        return 'мокрота и т.д.', 'мокрота'
    elif 'кровь' in material_lower:
        return 'кровь', 'катетер и кровь'
    elif 'катетер' in material_lower:
        return 'катетер', 'катетер и кровь'
    elif 'ликвор' in material_lower or 'цереброспинальная' in material_lower:
        return 'жидкость цереброспинальная', 'церебр'
    elif any(x in material_lower for x in ['рана', 'язва', 'соскоб', 'гной', 'абсцесс', 'отделяемое', 'эрозий', 'раневое', 'содержимое', 'пунктат']):
        return 'РАНЫ', 'раны'
    else:
        return 'другие', 'другие'

# ----------------------------------------------------------------------
# 3. Присвоение порядкового номера забора для каждого пациента и материала
# ----------------------------------------------------------------------

def assign_order(df):
    df = df.sort_values(['patient_id', 'material_raw', 'date']).reset_index(drop=True)
    order = []
    for (pid, mat), group in df.groupby(['patient_id', 'material_raw']):
        group = group.sort_values('date')
        orders = range(1, len(group)+1)
        order.extend(orders)
    df['order'] = order
    return df

# ----------------------------------------------------------------------
# 4. Вспомогательная функция подсчёта статистики
# ----------------------------------------------------------------------

def compute_stats(subdf):
    if subdf.empty:
        return (0,0,0,0,0.0,0,0.0)
    patients = subdf['patient_id'].nunique()
    total = len(subdf)
    no_growth = (~subdf['is_positive']).sum()
    pos = subdf['is_positive'].sum()
    pct_pos = (pos / total * 100) if total > 0 else 0.0
    diag = subdf[subdf['is_positive'] & subdf['diagnostic']].shape[0]
    pct_diag = (diag / pos * 100) if pos > 0 else 0.0
    return patients, total, no_growth, pos, pct_pos, diag, pct_diag

# ----------------------------------------------------------------------
# 5. Формирование листа 2 (Анализ результатов АБЧ общий)
# ----------------------------------------------------------------------

def generate_sheet2(df, log_func=None):
    df['cat2'] = df['material_raw'].apply(lambda x: categorize_material(x)[0])
    categories = ['всего', 'моча', 'мокрота и т.д.', 'кровь', 'катетер',
                  'жидкость цереброспинальная', 'РАНЫ', 'другие']
    rows = []
    for cat in categories:
        if log_func:
            log_func(f"Обработка категории: {cat}")
        if cat == 'всего':
            subdf = df
        else:
            subdf = df[df['cat2'] == cat]
        pat1, tot1, nog1, pos1, pct1, diag1, pctd1 = compute_stats(subdf)
        temp = subdf.copy()
        temp = assign_order(temp)
        subdf_ge2 = temp[temp['order'] >= 2]
        pat2, tot2, nog2, pos2, pct2, diag2, pctd2 = compute_stats(subdf_ge2)
        subdf_ge3 = temp[temp['order'] >= 3]
        pat3, tot3, nog3, pos3, pct3, diag3, pctd3 = compute_stats(subdf_ge3)
        rows.append({
            'Категория': cat,
            'кол-во больных': pat1, 'всего проб': tot1, 'роста нет': nog1, '+': pos1,
            '%': pct1, 'в диагностическом титре': diag1, '%_диаг': pctd1,
            'кол-во больных_2': pat2, 'всего проб_2': tot2, 'роста нет_2': nog2, '+_2': pos2,
            '%_2': pct2, 'в диагностическом титре_2': diag2, '%_диаг_2': pctd2,
            'кол-во больных_3': pat3, 'всего проб_3': tot3, 'роста нет_3': nog3, '+_3': pos3,
            '%_3': pct3, 'в диагностическом титре_3': diag3, '%_диаг_3': pctd3,
        })
    result_df = pd.DataFrame(rows)
    result_df = result_df.rename(columns={
        '%_диаг': '% (в диагн. титре)',
        '%_диаг_2': '% (в диагн. титре)',
        '%_диаг_3': '% (в диагн. титре)',
        '%': '% (+)', '%_2': '% (+)', '%_3': '% (+)'
    })
    return result_df

# ----------------------------------------------------------------------
# 6. Формирование листа 3 (Эпидемиология по локализациям)
# ----------------------------------------------------------------------

def generate_sheet3(df, log_func=None):
    df['cat3'] = df['material_raw'].apply(lambda x: categorize_material(x)[1])
    pos_df = df[df['is_positive']].copy()
    if pos_df.empty:
        if log_func:
            log_func("Нет положительных проб для листа 3.")
        return {}
    pos_df = assign_order(pos_df)
    organisms = pos_df['organism'].dropna().unique()
    categories3 = ['моча', 'мокрота', 'другие', 'церебр', 'катетер и кровь', 'раны']

    def build_table_for_category(cat):
        if cat == 'ВСЕ':
            subdf = pos_df
        else:
            subdf = pos_df[pos_df['cat3'] == cat]
        if subdf.empty:
            return None
        data = []
        for org in sorted(organisms):
            org_df = subdf[subdf['organism'] == org]
            if org_df.empty:
                continue
            row = {'Микроорганизм': org}
            for order in [1,2,3]:
                ord_df = org_df[org_df['order'] == order]
                total = len(ord_df)
                if total == 0:
                    row[f'diag_{order}'] = 0
                    row[f'nediag_{order}'] = 0
                    row[f'total_{order}'] = 0
                    row[f'pct_{order}'] = 0.0
                else:
                    diag = ord_df['diagnostic'].sum()
                    non_diag = total - diag
                    pct = (diag / total * 100) if total > 0 else 0.0
                    row[f'diag_{order}'] = diag
                    row[f'nediag_{order}'] = non_diag
                    row[f'total_{order}'] = total
                    row[f'pct_{order}'] = pct
            data.append(row)
        total_row = {'Микроорганизм': 'ВСЕ'}
        for order in [1,2,3]:
            total_diag = subdf[subdf['order'] == order]['diagnostic'].sum()
            total_all = subdf[subdf['order'] == order].shape[0]
            total_non = total_all - total_diag
            pct = (total_diag / total_all * 100) if total_all > 0 else 0.0
            total_row[f'diag_{order}'] = total_diag
            total_row[f'nediag_{order}'] = total_non
            total_row[f'total_{order}'] = total_all
            total_row[f'pct_{order}'] = pct
        data.append(total_row)
        df_cat = pd.DataFrame(data)
        col_map = {}
        for order in [1,2,3]:
            col_map[f'diag_{order}'] = f'диаг_{order}'
            col_map[f'nediag_{order}'] = f'недиаг_{order}'
            col_map[f'total_{order}'] = f'итого_{order}'
            col_map[f'pct_{order}'] = f'% диагн_{order}'
        df_cat = df_cat.rename(columns=col_map)
        return df_cat

    result = {}
    for cat in categories3:
        if log_func:
            log_func(f"Обработка категории листа 3: {cat}")
        tbl = build_table_for_category(cat)
        if tbl is not None:
            result[cat] = tbl
    total_tbl = build_table_for_category('ВСЕ')
    if total_tbl is not None:
        result['ВСЕ'] = total_tbl
    return result

# ----------------------------------------------------------------------
# 7. Формирование листа 4 (Антибиотикорезистентность) – общая
#    и листа 5 (раневые антибиотики с локализацией)
# ----------------------------------------------------------------------

def generate_antibiotic_sheet(ab_data, log_func=None):
    """
    Создаёт DataFrame для листа с антибиотикорезистентностью.
    Для каждого антибиотика выводит таблицу с результатами.
    Если индекс DataFrame содержит несколько уровней (например, organism и material_raw),
    то будет добавлена колонка "Локализация".
    """
    if not ab_data:
        if log_func:
            log_func("Нет данных по антибиотикам для листа")
        return pd.DataFrame()

    all_rows = []
    for antibiotic, df_ab in ab_data.items():
        if df_ab.empty:
            continue
        # Определяем, является ли индекс MultiIndex
        is_multi = isinstance(df_ab.index, pd.MultiIndex)
        # Заголовок антибиотика
        all_rows.append([f"Антибиотик: {antibiotic}"])
        # Заголовки столбцов
        if is_multi:
            header = ['Микроорганизм', 'Локализация', 'S', 'I', 'R', 'НД', 'Total', 'S%', 'I%', 'R%', 'НД%']
        else:
            header = ['Микроорганизм', 'S', 'I', 'R', 'НД', 'Total', 'S%', 'I%', 'R%', 'НД%']
        all_rows.append(header)

        # Обходим строки
        for idx, row in df_ab.iterrows():
            if is_multi:
                # idx - кортеж (organism, material_raw)
                organism, location = idx
                row_data = [
                    organism,
                    location,
                    row['S'], row['I'], row['R'], row['НД'],
                    row['Total'],
                    round(row['S_%'], 1), round(row['I_%'], 1),
                    round(row['R_%'], 1), round(row['НД_%'], 1)
                ]
            else:
                organism = idx
                row_data = [
                    organism,
                    row['S'], row['I'], row['R'], row['НД'],
                    row['Total'],
                    round(row['S_%'], 1), round(row['I_%'], 1),
                    round(row['R_%'], 1), round(row['НД_%'], 1)
                ]
            all_rows.append(row_data)
        all_rows.append([])  # пустая строка после таблицы
    # Убираем последнюю пустую строку
    if all_rows and all_rows[-1] == []:
        all_rows.pop()
    # Преобразуем в DataFrame
    df_result = pd.DataFrame(all_rows)
    return df_result

# (для обратной совместимости оставляем псевдоним)
generate_sheet4 = generate_antibiotic_sheet


def build_wound_sheet5(data, antibiotics, ab_indices, log_func=None):
    """Формирует лист 5: раневые антибиотики с разбивкой по локализации."""
    cat_wound = data['material_raw'].apply(lambda x: categorize_material(x)[0])
    wound_data = data[cat_wound == 'РАНЫ'].copy()
    if wound_data.empty:
        if log_func:
            log_func("Нет раневых материалов, лист 5 будет пуст.")
        return pd.DataFrame(), wound_data
    if log_func:
        log_func(f"Раневых образцов: {len(wound_data)}")
    ab_data_wound = compute_ab_data(
        wound_data, antibiotics, ab_indices,
        group_cols=['organism', 'material_raw'],
        log_func=log_func
    )
    sheet5 = generate_antibiotic_sheet(ab_data_wound, log_func=log_func)
    return sheet5, wound_data


def run_pipeline(input_path, output_path, styled=True, log_func=None):
    """Полный цикл обработки: загрузка -> листы 2..5 -> сохранение в Excel."""
    df_main, ab_data, data, antibiotics, ab_indices = load_data(input_path, log_func=log_func)

    if log_func:
        log_func("Формирование листа 2 (анализ посевов по категориям)...")
    sheet2 = generate_sheet2(df_main, log_func=log_func)

    if log_func:
        log_func("Формирование листа 3 (эпидемиология по локализациям)...")
    sheet3 = generate_sheet3(df_main, log_func=log_func)

    if log_func:
        log_func("Формирование листа 4 (общая антибиотикорезистентность)...")
    sheet4 = generate_antibiotic_sheet(ab_data, log_func=log_func)

    if log_func:
        log_func("Формирование листа 5 (раневые антибиотики с локализацией)...")
    sheet5, wound_data = build_wound_sheet5(data, antibiotics, ab_indices, log_func=log_func)

    save_to_excel(sheet2, sheet3, sheet4, sheet5, output_path,
                  log_func=log_func, styled=styled,
                  raw_data=data, wound_data=wound_data)

    if log_func:
        log_func("Обработка успешно завершена!")
    return dict(df_main=df_main, ab_data=ab_data, sheet2=sheet2, sheet3=sheet3,
                sheet4=sheet4, sheet5=sheet5)

# ----------------------------------------------------------------------
# 8. Сохранение в Excel
# ----------------------------------------------------------------------

def save_to_excel(sheet2_df, sheet3_dict, sheet4_df, sheet5_df, output_file,
                  log_func=None, styled=False, raw_data=None, wound_data=None):
    """
    Сохраняет аналитические листы в Excel.
    styled=True — современное оформление (шапки, зебра, заморозка панелей,
    условное форматирование % чувствительности) + дополнительные листы
    «Сводка» и «Данные».
    """
    if log_func:
        log_func(f"Сохранение в {output_file}...")
    dirname = os.path.dirname(os.path.abspath(output_file))
    if dirname:
        os.makedirs(dirname, exist_ok=True)

    with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
        sheet2_df.to_excel(writer, sheet_name='2', index=False)

        # Лист 3
        sheet3_rows = []
        for section_name, section_df in sheet3_dict.items():
            sheet3_rows.append([section_name] + [''] * (len(section_df.columns)-1))
            sheet3_rows.append(section_df.columns.tolist())
            for _, row in section_df.iterrows():
                sheet3_rows.append(row.tolist())
            sheet3_rows.append([''] * len(section_df.columns))
        if sheet3_rows:
            max_cols = max(len(row) for row in sheet3_rows)
            for row in sheet3_rows:
                if len(row) < max_cols:
                    row.extend([''] * (max_cols - len(row)))
            sheet3_combined = pd.DataFrame(sheet3_rows)
            sheet3_combined.to_excel(writer, sheet_name='3', index=False, header=False)

        # Лист 4
        if not sheet4_df.empty:
            sheet4_df.to_excel(writer, sheet_name='4', index=False, header=False)

        # Лист 5 (раневые антибиотики с локализацией)
        if not sheet5_df.empty:
            sheet5_df.to_excel(writer, sheet_name='5', index=False, header=False)

        if styled:
            extra_sheets = []
            summary_df = _build_summary(raw_data, sheet2_df)
            if summary_df is not None:
                summary_df.to_excel(writer, sheet_name='Сводка', index=False)
                extra_sheets.append('Сводка')
            if raw_data is not None and not raw_data.empty:
                _write_raw_sheet(raw_data, writer)
                extra_sheets.append('Данные')
            _apply_styling(writer, sheet2_df, sheet3_dict, sheet4_df, sheet5_df, extra_sheets)

        else:
            # Автонастройка ширины столбцов (как в базовой версии)
            for sheetname in writer.sheets:
                worksheet = writer.sheets[sheetname]
                for column in worksheet.columns:
                    max_length = 0
                    column_letter = get_column_letter(column[0].column)
                    for cell in column:
                        try:
                            if cell.value:
                                max_length = max(max_length, len(str(cell.value)))
                        except Exception:
                            pass
                    adjusted_width = min(max_length + 2, 50)
                    worksheet.column_dimensions[column_letter].width = adjusted_width

    if log_func:
        log_func("Сохранение завершено.")


def _build_summary(raw_data, sheet2_df):
    """Небольшая сводная таблица ключевых показателей для листа «Сводка»."""
    if raw_data is None or getattr(raw_data, 'empty', True):
        return None
    total = len(raw_data)
    pos = int(raw_data['is_positive'].sum())
    diag = int((raw_data['is_positive'] & raw_data['diagnostic']).sum())
    patients = raw_data['patient_id'].nunique()
    dates = raw_data['date'].dropna()
    period = ''
    if not dates.empty:
        period = f"{dates.min():%d.%m.%Y} — {dates.max():%d.%m.%Y}"
    top_org = ''
    orgs = raw_data.loc[raw_data['is_positive'], 'organism'].dropna()
    if not orgs.empty:
        top = orgs.value_counts().head(3)
        top_org = ', '.join(f"{name} ({cnt})" for name, cnt in top.items())
    rows = [
        ('Период', period),
        ('Пациентов', patients),
        ('Всего проб', total),
        ('Положительных проб', pos),
        ('Доля положительных, %', round(pos / total * 100, 1) if total else 0),
        ('В диагностическом титре', diag),
        ('Доля диагностичных, %', round(diag / pos * 100, 1) if pos else 0),
        ('Лидеры по выделению', top_org),
    ]
    return pd.DataFrame(rows, columns=['Показатель', 'Значение'])


def _write_raw_sheet(raw_data, writer):
    """Лист «Данные»: очищенные строки посевов с категорией материала."""
    df = raw_data.copy()
    cats = df['material_raw'].apply(lambda x: categorize_material(x)[0])
    out = pd.DataFrame({
        'Пациент': df['patient_id'],
        'Дата': df['date'].dt.strftime('%d.%m.%Y'),
        'Материал': df['material_raw'],
        'Категория': cats,
        'Микроорганизм': df['organism'],
        'Рост': df['is_positive'].map({True: '+', False: 'нет'}),
        'Диагностичный': df['diagnostic'].map({True: 'да', False: 'нет'}),
    })
    out.to_excel(writer, sheet_name='Данные', index=False)


def _apply_styling(writer, sheet2_df, sheet3_dict, sheet4_df, sheet5_df, extra_sheets):
    """Современное оформление книги: шапки, зебра, числовые форматы, панели."""
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.utils import get_column_letter as gcl

    ACCENT = "1F4E79"      # тёмно-синий (шапки)
    ACCENT_SOFT = "DCE6F1"  # светло-голубой (зебра/секции)
    ZEBRA = "F5F8FC"
    GOOD = "C6EFCE"
    BAD = "FFC7CE"

    thin = Side(style="thin", color="B8C4D4")
    border_all = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor=ACCENT)
    zebra_fill = PatternFill("solid", fgColor=ZEBRA)
    section_fill = PatternFill("solid", fgColor=ACCENT_SOFT)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_al = Alignment(horizontal="left", vertical="center")

    def style_header_row(ws, row_idx, ncols):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=row_idx, column=c)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            cell.border = border_all

    def autofit(ws, max_width=50):
        widths = {}
        for row in ws.iter_rows(values_only=False):
            for cell in row:
                if cell.value is not None:
                    l = len(str(cell.value))
                    widths[cell.column] = max(widths.get(cell.column, 0), l)
        for col, length in widths.items():
            ws.column_dimensions[gcl(col)].width = min(length + 2.5, max_width)

    # --- Сводка / Данные / лист 2: простые таблицы с шапкой -------------
    simple_with_header = [s for s in ['Сводка', 'Данные', '2'] if s in writer.sheets]
    for name in simple_with_header:
        ws = writer.sheets[name]
        ncols = ws.max_column
        style_header_row(ws, 1, ncols)
        for r in range(2, ws.max_row + 1):
            for c in range(1, ncols + 1):
                cell = ws.cell(row=r, column=c)
                cell.border = border_all
                if r % 2 == 0:
                    cell.fill = zebra_fill
        ws.freeze_panes = "A2"
        autofit(ws)

    # --- Лист 3: секции (заголовок + шапка + строки) ---------------------
    if '3' in writer.sheets and sheet3_dict:
        ws = writer.sheets['3']
        ncols = ws.max_column
        r = 1
        max_row = ws.max_row
        while r <= max_row:
            a = ws.cell(row=r, column=1).value
            if a in (None, ''):
                r += 1
                continue
            b = ws.cell(row=r + 1, column=1).value if r + 1 <= max_row else None
            if b == 'Микроорганизм':
                style_header_row(ws, r + 1, ncols)
                rr = r + 2
                k = 0
                while rr <= max_row:
                    val = ws.cell(row=rr, column=1).value
                    if val in (None, ''):
                        break
                    for c in range(1, ncols + 1):
                        cell = ws.cell(row=rr, column=c)
                        cell.border = border_all
                        if k % 2 == 1:
                            cell.fill = zebra_fill
                    if str(val) == 'ВСЕ':
                        for c in range(1, ncols + 1):
                            ws.cell(row=rr, column=c).font = Font(bold=True)
                    k += 1
                    rr += 1
                # подсветка процентов диагностичности
                for c in range(1, ncols + 1):
                    head = ws.cell(row=r + 1, column=c).value
                    if head and str(head).startswith('%'):
                        rng = f"{gcl(c)}{r + 2}:{gcl(c)}{r + 1 + k}"
                        ws.conditional_formatting.add(rng, CellIsRule(
                            operator='greaterThanOrEqual', formula=['80'],
                            fill=PatternFill(start_color=GOOD, end_color=GOOD, fill_type='solid')))
                        ws.conditional_formatting.add(rng, CellIsRule(
                            operator='lessThan', formula=['50'],
                            fill=PatternFill(start_color=BAD, end_color=BAD, fill_type='solid')))
                ws.cell(row=r, column=1).fill = section_fill
                ws.cell(row=r, column=1).font = Font(bold=True, color=ACCENT, size=12)
                r = rr + 1
            else:
                r += 1
        autofit(ws)

    # --- Листы 4 и 5: блоки «Антибиотик: ...» ----------------------------
    for name, block_df in (('4', sheet4_df), ('5', sheet5_df)):
        if name not in writer.sheets or block_df is None or block_df.empty:
            continue
        ws = writer.sheets[name]
        ncols = ws.max_column
        r = 1
        max_row = ws.max_row
        while r <= max_row:
            val = ws.cell(row=r, column=1).value
            if val and str(val).startswith('Антибиотик:'):
                ws.cell(row=r, column=1).font = Font(bold=True, color=ACCENT, size=12)
                ws.cell(row=r, column=1).fill = section_fill
                if r + 1 <= max_row:
                    style_header_row(ws, r + 1, ncols)
                rr = r + 2
                k = 0
                s_col = None
                while rr <= max_row:
                    v = ws.cell(row=rr, column=1).value
                    if v in (None, '') or str(v).startswith('Антибиотик:'):
                        break
                    for c in range(1, ncols + 1):
                        cell = ws.cell(row=rr, column=c)
                        cell.border = border_all
                        if k % 2 == 1:
                            cell.fill = zebra_fill
                    if s_col is None:
                        for c in range(1, ncols + 1):
                            if ws.cell(row=r + 1, column=c).value == 'S%':
                                s_col = c
                                break
                    k += 1
                    rr += 1
                if s_col:
                    rng = f"{gcl(s_col)}{r + 2}:{gcl(s_col)}{r + 1 + k}"
                    ws.conditional_formatting.add(rng, CellIsRule(
                        operator='greaterThanOrEqual', formula=['80'],
                        fill=PatternFill(start_color=GOOD, end_color=GOOD, fill_type='solid')))
                    ws.conditional_formatting.add(rng, CellIsRule(
                        operator='lessThan', formula=['50'],
                        fill=PatternFill(start_color=BAD, end_color=BAD, fill_type='solid')))
                r = rr
            else:
                r += 1
        autofit(ws)

    # Активный лист — «Сводка», если он есть
    if 'Сводка' in writer.sheets:
        writer.book.active = writer.sheets['Сводка']
